
import argparse
import asyncio
import logging
import sys
from pathlib import Path

import asyncpg
import yaml
from openpyxl import load_workbook
for stream in (sys.stdout, sys.stderr):
    if hasattr(stream, "reconfigure"):
        stream.reconfigure(encoding="utf-8", errors="replace")

logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
logger = logging.getLogger(__name__)

SHEET_NAME = "операции+сроки"
TABLE_NAME = "static_initial_data_control_operations"
HEADER_ROW = 1
DATA_START_ROW = 3
REGION_COLUMNS = {
    "regionbels": "Белгород Юг",
    "regionbelc": "Белгород Центр",
    "regionbelзапад": "Белгород Запад",
    "regionbelk": "Белгород-Курск",
    "regiontams": "Тамбов-Юг",
    "regiontamn": "Тамбов-Север",
    "regionorel": "Орел",
    "regionsara": "Саратов НПК",
    "regionprim": "Приморье",
}

REQUIRED_COLUMNS = [
    "culture_id",
    "t_material",
    "season",
    "input_operation_order",
    "input_deadline",
    "noinput_deadline",
    "order",
    "year",
]


class LoadError(Exception):
    """Ошибка, из-за которой загружать данные нельзя."""

def normalize_list_cell(value):
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    parts = [p.strip() for p in text.split(",")]
    parts = [p for p in parts if p]
    if not parts:
        return None
    return ", ".join(parts)


def parse_int_list(text):
    if text is None:
        return []
    out = []
    for token in str(text).split(","):
        token = token.strip()
        if not token:
            continue
        try:
            out.append(int(token))
        except ValueError as exc:
            raise LoadError(f"нечисловое значение в списке: '{token}'") from exc
    return out


def to_date(value):
    if value is None:
        return None
    if hasattr(value, "date"):
        return value.date()
    if isinstance(value, str) and not value.strip():
        return None
    raise LoadError(f"ожидалась дата, получено {value!r}")


def read_excel(path):
    """Читает лист операций и возвращает список словарей """
    workbook = load_workbook(path, data_only=True, read_only=True)
    if SHEET_NAME not in workbook.sheetnames:
        raise LoadError(
            f"в файле нет листа '{SHEET_NAME}'. Найдено: {', '.join(workbook.sheetnames)}"
        )
    sheet = workbook[SHEET_NAME]

    rows = list(sheet.iter_rows(values_only=True))
    if len(rows) < DATA_START_ROW:
        raise LoadError("лист пустой или короче двух строк заголовков")

    header = [str(c).strip() if c is not None else "" for c in rows[HEADER_ROW - 1]]
    index_of = {name: i for i, name in enumerate(header) if name}

    missing = [c for c in REQUIRED_COLUMNS if c not in index_of]
    if missing:
        raise LoadError(f"в листе нет обязательных колонок: {', '.join(missing)}")

    missing_regions = [c for c in REGION_COLUMNS if c not in index_of]
    if missing_regions:
        raise LoadError(f"в листе нет колонок регионов: {', '.join(missing_regions)}")

    def cell(row, column):
        position = index_of[column]
        return row[position] if position < len(row) else None

    records = []
    for row_number, row in enumerate(rows[DATA_START_ROW - 1:], start=DATA_START_ROW):
        if cell(row, "culture_id") in (None, ""):
            continue  # хвостовые пустые строки

        orders_text = normalize_list_cell(cell(row, "input_operation_order"))
        deadlines_text = normalize_list_cell(cell(row, "input_deadline"))
        orders = parse_int_list(orders_text)
        deadlines = parse_int_list(deadlines_text)
        if len(orders) != len(deadlines):
            raise LoadError(
                f"строка {row_number}: не совпадает количество зависимостей и сроков — "
                f"'{orders_text}' против '{deadlines_text}'"
            )

        noinput = cell(row, "noinput_deadline")
        records.append(
            {
                "excel_row": row_number,
                "culture_id": int(cell(row, "culture_id")),
                "t_material": str(cell(row, "t_material")).strip(),
                "season": str(cell(row, "season")).strip(),
                "input_operation_order": orders_text,
                "input_deadline": deadlines_text,
                "noinput_deadline": int(noinput) if noinput not in (None, "") else None,
                "order": int(cell(row, "order")),
                "year": int(cell(row, "year")),
                "region_dates": {
                    column: to_date(cell(row, column)) for column in REGION_COLUMNS
                },
            }
        )

    workbook.close()
    return records



def drop_duplicates(records, strict):
    seen_order, seen_material = {}, {}
    kept, dropped = [], []

    for record in records:
        order_key = (record["culture_id"], record["order"], record["year"])
        material_key = (
            record["culture_id"],
            record["t_material"],
            record["season"],
            record["year"],
        )

        clash = seen_order.get(order_key) or seen_material.get(material_key)
        if clash:
            dropped.append((record, clash))
            continue

        seen_order[order_key] = record
        seen_material[material_key] = record
        kept.append(record)

    if dropped:
        for record, first in dropped:
            logger.warning(
                "  строка %s: culture=%s t_material=%s season=%s order=%s "
                "— конфликтует со строкой %s (order=%s); строка ПРОПУЩЕНА",
                record["excel_row"], record["culture_id"], record["t_material"],
                record["season"], record["order"], first["excel_row"], first["order"],
            )
        if strict:
            raise LoadError(
                "загрузка прервана из-за дубликатов (--strict). "
            )

    return kept


def resolve_ids(records, material_ids, region_ids):
    unknown_materials = sorted(
        {r["t_material"] for r in records if r["t_material"] not in material_ids}
    )
    if unknown_materials:
        raise LoadError(
            f"в справочнике sap_operations нет операций ({len(unknown_materials)}): "
            + ", ".join(unknown_materials[:10])
            + (" ..." if len(unknown_materials) > 10 else "")
        )

    unknown_regions = sorted(set(REGION_COLUMNS.values()) - set(region_ids))
    if unknown_regions:
        raise LoadError(
            "в справочнике sap_regions нет регионов: " + ", ".join(unknown_regions)
        )

    rows = []
    for record in records:
        for column, region_name in REGION_COLUMNS.items():
            rows.append(
                (
                    record["culture_id"],
                    material_ids[record["t_material"]],
                    region_ids[region_name],
                    record["season"],
                    record["region_dates"][column],
                    record["input_operation_order"],
                    record["input_deadline"],
                    record["noinput_deadline"],
                    record["order"],
                    record["year"],
                )
            )
    return rows


def report_dependency_stats(records):
    histogram = {}
    for record in records:
        count = len(parse_int_list(record["input_operation_order"]))
        histogram[count] = histogram.get(count, 0) + 1
    for count in sorted(histogram):
        label = "без зависимостей" if count == 0 else f"{count} зависимост(ь/и)"
        logger.info("  %-22s %d строк", label, histogram[count])


def load_db_config(path):
    with open(path, "r", encoding="utf-8") as handle:
        config = yaml.safe_load(handle)
    required = ["host", "port", "user", "password", "dbname"]
    missing = [key for key in required if key not in config]
    if missing:
        raise LoadError(f"в {path} нет полей: {', '.join(missing)}")
    return config


async def fetch_lookup(connection, query, key_field, value_field):
    rows = await connection.fetch(query)
    return {row[key_field]: row[value_field] for row in rows}


async def write_rows(connection, rows, years):
    """Заменяет записи только за те годы, которые есть в Excel."""
    async with connection.transaction():
        await connection.execute(
            f"DELETE FROM {TABLE_NAME} WHERE year = ANY($1::int[])", years
        )
        await connection.copy_records_to_table(
            TABLE_NAME,
            records=rows,
            columns=[
                "culture_id",
                "t_material_id",
                "region_id",
                "season",
                "region_date",
                "input_operation_order",
                "input_deadline",
                "noinput_deadline",
                "order",
                "year",
            ],
        )


async def run(args):
    excel_path = Path(args.excel)
    if not excel_path.exists():
        raise LoadError(f"файл не найден: {excel_path}")

    logger.info("Читаю %s", excel_path.name)
    records = read_excel(excel_path)
    logger.info("Строк с данными: %d", len(records))
    report_dependency_stats(records)

    records = drop_duplicates(records, strict=args.strict)

    config = load_db_config(args.config)
    connection = await asyncpg.connect(
        host=config["host"], port=config["port"], user=config["user"],
        password=config["password"], database=config["dbname"],
    )
    try:
        material_ids = await fetch_lookup(
            connection, "SELECT t_material, id FROM sap_operations", "t_material", "id"
        )
        region_ids = await fetch_lookup(
            connection, "SELECT name, id FROM sap_regions", "name", "id"
        )
        logger.info(
            "Справочники: операций %d, регионов %d", len(material_ids), len(region_ids)
        )

        rows = resolve_ids(records, material_ids, region_ids)
        logger.info(
            "К загрузке: %d строк (%d операций x %d регионов)",
            len(rows), len(records), len(REGION_COLUMNS),
        )

        years = sorted({record["year"] for record in records})

        if args.dry_run:
            logger.info("--dry-run: база не изменена")
            return

        await write_rows(connection, rows, years)
        logger.info(
            "Готово: в %s загружено %d строк за %s",
            TABLE_NAME, len(rows), ", ".join(str(y) for y in years),
        )
    finally:
        await connection.close()


def main():
    parser = argparse.ArgumentParser(
        description="Загрузка справочника операций из Excel в static_initial_data_control_operations"
    )
    parser.add_argument("excel", help="путь к файлу 'инит дата ... .xlsx'")
    parser.add_argument(
        "--config", default="config/db_config.yml", help="конфиг БД (по умолчанию config/db_config.yml)"
    )
    parser.add_argument(
        "--dry-run", action="store_true", help="только проверить файл, ничего не записывать"
    )
    parser.add_argument(
        "--strict", action="store_true", help="прерваться при дубликатах вместо пропуска"
    )
    args = parser.parse_args()

    try:
        asyncio.run(run(args))
    except LoadError as error:
        logger.error("%s", error)
        sys.exit(1)


if __name__ == "__main__":
    main()
